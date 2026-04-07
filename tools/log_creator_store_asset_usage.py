from __future__ import annotations

import argparse
import json
import re
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
RETENTION_LIMIT = 5


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Write a rolling Excel log of Creator Store assets used in the current Studio world."
    )
    parser.add_argument(
        "--experience-root",
        required=True,
        help="Experience root folder that contains prompt_world_experience.json.",
    )
    parser.add_argument(
        "--log-root",
        default="docs/creator_store_asset_logs",
        help="Folder where the rolling Excel workbooks will be written.",
    )
    return parser.parse_args()


def slugify(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", value.strip()).strip("-")
    return cleaned or "world"


def load_experience_metadata(experience_root: Path) -> dict[str, Any]:
    metadata_path = experience_root / "prompt_world_experience.json"
    return json.loads(metadata_path.read_text(encoding="utf-8"))


def write_experience_metadata(experience_root: Path, payload: dict[str, Any]) -> None:
    metadata_path = experience_root / "prompt_world_experience.json"
    metadata_path.write_text(json.dumps(payload, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def run_bridge_code(repo_root: Path, command: str) -> str:
    result = subprocess.run(
        [
            str(repo_root / "integrations" / "roblox-studio-mcp" / "studio-bridge.cmd"),
            "run-code",
            command,
        ],
        cwd=repo_root,
        check=True,
        capture_output=True,
        text=True,
    )

    output = result.stdout.strip()
    if output.startswith("[OUTPUT]"):
        output = output[len("[OUTPUT]") :].strip()
    return output


def fetch_live_asset_usage(repo_root: Path) -> dict[str, Any]:
    # Keep this on one line because the Studio bridge is more reliable with compact Luau snippets.
    studio_command = (
        "local HttpService=game:GetService('HttpService'); "
        "local payload={"
        "title=workspace:GetAttribute('PromptWorldTitle'), "
        "generation_seed=workspace:GetAttribute('PromptWorldGenerationSeed'), "
        "originality_signature=workspace:GetAttribute('PromptWorldOriginalitySignature'), "
        "render_style=workspace:GetAttribute('PromptWorldRenderStyle'), "
        "terrain=workspace:GetAttribute('PromptWorldTerrain'), "
        "unique_asset_count=workspace:GetAttribute('PromptWorldCreatorStoreAssetCount'), "
        "placed_instance_count=workspace:GetAttribute('PromptWorldCreatorStorePlacedInstanceCount'), "
        "asset_usage_json=workspace:GetAttribute('PromptWorldCreatorStoreAssetUsage')"
        "}; "
        "print(HttpService:JSONEncode(payload))"
    )

    output = run_bridge_code(repo_root, studio_command)
    payload = json.loads(output)
    asset_usage_json = payload.get("asset_usage_json")
    payload["asset_usage"] = json.loads(asset_usage_json) if asset_usage_json else []
    return payload


def load_manifest(log_root: Path) -> list[dict[str, Any]]:
    manifest_path = log_root / "asset_log_manifest.json"
    if not manifest_path.exists():
        return []
    return json.loads(manifest_path.read_text(encoding="utf-8"))


def write_manifest(log_root: Path, entries: list[dict[str, Any]]) -> None:
    manifest_path = log_root / "asset_log_manifest.json"
    manifest_path.write_text(json.dumps(entries, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def style_headers(sheet) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def auto_fit_columns(sheet) -> None:
    for column_cells in sheet.columns:
        values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
        width = min(max(len(value) for value in values) + 2, 48)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def build_workbook(
    *,
    workbook_path: Path,
    log_entry: dict[str, Any],
    experience_metadata: dict[str, Any],
    live_payload: dict[str, Any],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    style_headers(summary)

    summary_rows = [
        ("World Log Id", log_entry["world_log_id"]),
        ("Created At UTC", log_entry["created_at_utc"]),
        ("Experience Name", experience_metadata["name"]),
        ("World Title", live_payload.get("title")),
        ("World Prompt", experience_metadata.get("world_prompt")),
        ("Generation Seed", live_payload.get("generation_seed") or experience_metadata.get("generation_seed")),
        ("Originality Signature", live_payload.get("originality_signature")),
        ("Render Style", live_payload.get("render_style")),
        ("Terrain", live_payload.get("terrain")),
        ("Unique Creator Store Assets", live_payload.get("unique_asset_count") or len(live_payload["asset_usage"])),
        ("Placed Asset Instances", live_payload.get("placed_instance_count") or 0),
        ("Workbook Retention Limit", RETENTION_LIMIT),
    ]
    for row in summary_rows:
        summary.append(row)

    assets_sheet = workbook.create_sheet("Assets")
    assets_sheet.append(
        [
            "World Log Id",
            "Asset Sequence",
            "Identification Number",
            "Inserted Model Name",
            "Marketplace Name",
            "Search Query",
            "Asset Id",
            "Asset Version Id",
            "Creator Name",
            "Library URL",
            "Asset Role",
            "Placement Count",
        ]
    )
    style_headers(assets_sheet)

    total_placements = 0
    for index, asset in enumerate(live_payload["asset_usage"], start=1):
        placement_count = int(asset.get("placement_count") or 0)
        total_placements += placement_count
        assets_sheet.append(
            [
                log_entry["world_log_id"],
                index,
                asset.get("identification_number") or asset.get("asset_id") or "",
                asset.get("inserted_model_name") or "",
                asset.get("marketplace_name") or "",
                asset.get("query") or "",
                asset.get("asset_id") or "",
                asset.get("asset_version_id") or "",
                asset.get("creator_name") or "",
                asset.get("library_url") or "",
                asset.get("asset_role") or "",
                placement_count,
            ]
        )

    counts_sheet = workbook.create_sheet("Counts")
    counts_sheet.append(["Metric", "Value"])
    style_headers(counts_sheet)
    counts_sheet.append(["Unique Assets Used", len(live_payload["asset_usage"])])
    counts_sheet.append(["Placed Asset Instances", total_placements])

    auto_fit_columns(summary)
    auto_fit_columns(assets_sheet)
    auto_fit_columns(counts_sheet)
    workbook.save(workbook_path)


def enforce_retention(log_root: Path, entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    sorted_entries = sorted(entries, key=lambda entry: entry["world_log_id"])
    while len(sorted_entries) > RETENTION_LIMIT:
        removed = sorted_entries.pop(0)
        workbook_path = log_root / removed["workbook_name"]
        if workbook_path.exists():
            workbook_path.unlink()
    return sorted_entries


def update_experience_metadata(
    *,
    log_root: Path,
    experience_root: Path,
    experience_metadata: dict[str, Any],
    log_entry: dict[str, Any],
    manifest: list[dict[str, Any]],
) -> dict[str, Any]:
    workbook_path = (log_root / log_entry["workbook_name"]).resolve()
    retained_logs = [
        {
            "world_log_id": entry["world_log_id"],
            "workbook_name": entry["workbook_name"],
            "created_at_utc": entry["created_at_utc"],
        }
        for entry in manifest
    ]
    experience_metadata["asset_log"] = {
        "latest_world_log_id": log_entry["world_log_id"],
        "latest_workbook": str(workbook_path),
        "latest_created_at_utc": log_entry["created_at_utc"],
        "retention_limit": RETENTION_LIMIT,
        "unique_asset_count": log_entry["unique_asset_count"],
        "placed_instance_count": log_entry["placed_instance_count"],
        "retained_workbooks": retained_logs,
    }
    write_experience_metadata(experience_root, experience_metadata)
    return experience_metadata


def main() -> int:
    args = parse_args()
    repo_root = Path(__file__).resolve().parent.parent
    experience_root = Path(args.experience_root).resolve()
    log_root = (repo_root / args.log_root).resolve()
    log_root.mkdir(parents=True, exist_ok=True)

    experience_metadata = load_experience_metadata(experience_root)
    live_payload = fetch_live_asset_usage(repo_root)
    if not live_payload["asset_usage"]:
        raise SystemExit("No Creator Store asset usage is currently published in Studio.")

    manifest = load_manifest(log_root)
    next_log_id = (max((entry["world_log_id"] for entry in manifest), default=0) + 1)
    created_at = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H-%M-%SZ")
    workbook_name = (
        f"creator_store_assets_{next_log_id:04d}_{slugify(experience_metadata['name'])}_"
        f"{live_payload.get('originality_signature') or 'world'}.xlsx"
    )
    workbook_path = log_root / workbook_name

    log_entry = {
        "world_log_id": next_log_id,
        "created_at_utc": created_at,
        "experience_name": experience_metadata["name"],
        "world_title": live_payload.get("title"),
        "generation_seed": live_payload.get("generation_seed") or experience_metadata.get("generation_seed"),
        "originality_signature": live_payload.get("originality_signature"),
        "unique_asset_count": live_payload.get("unique_asset_count") or len(live_payload["asset_usage"]),
        "placed_instance_count": live_payload.get("placed_instance_count") or 0,
        "workbook_name": workbook_name,
    }

    build_workbook(
        workbook_path=workbook_path,
        log_entry=log_entry,
        experience_metadata=experience_metadata,
        live_payload=live_payload,
    )

    manifest.append(log_entry)
    manifest = enforce_retention(log_root, manifest)
    write_manifest(log_root, manifest)
    experience_metadata = update_experience_metadata(
        log_root=log_root,
        experience_root=experience_root,
        experience_metadata=experience_metadata,
        log_entry=log_entry,
        manifest=manifest,
    )

    print(
        json.dumps(
            {
                "workbook": str(workbook_path),
                "world_log_id": next_log_id,
                "retained_workbooks": [entry["workbook_name"] for entry in manifest],
                "retention_limit": RETENTION_LIMIT,
                "unique_asset_count": log_entry["unique_asset_count"],
                "placed_instance_count": log_entry["placed_instance_count"],
                "experience_metadata": str(experience_root / "prompt_world_experience.json"),
                "latest_experience_asset_log": experience_metadata["asset_log"],
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
